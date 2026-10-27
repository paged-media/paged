#!/usr/bin/env python3
"""Write the `charts` fixture's charts into PowerPoint's blank deck.

  python3 charts.seed.py <blank.pptx> <seed.pptx>

One chart per slide, each with its embedded workbook (written with
openpyxl) laid out as PowerPoint lays out a chart's sheet: a clustered
column chart, a line chart, a pie, a scatter chart, and a column chart
WITHOUT an embedded workbook (only the cached values in the chart part,
as a chart whose linked data is gone). PowerPoint then re-saves the deck,
so the chart parts and workbooks recorded are PowerPoint's."""
import io, re, sys, zipfile

import openpyxl

EMU = 12700
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

CATS = ["Q1", "Q2", "Q3", "Q4"]
SERIES = [("North", [4.3, 2.5, 3.5, 4.5]), ("South", [2.4, 4.4, 1.8, 2.8]), ("West", [2, 2, 3, 5])]
SCATTER_X = [1, 2, 3, 4, 5]
SCATTER_Y = [("Measured", [2.1, 3.9, 6.2, 7.8, 10.1])]

# (kind, series, embedded)
CHARTS = [
    ("bar", SERIES, True),
    ("line", SERIES, True),
    ("pie", SERIES[:1], True),
    ("scatter", SCATTER_Y, True),
    ("bar", SERIES[:2], False),
]


def col(i):
    return chr(ord("A") + i)


def workbook(kind, series):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    xs = SCATTER_X if kind == "scatter" else CATS
    ws.cell(1, 1, "X-Values" if kind == "scatter" else None)
    for i, (name, _) in enumerate(series):
        ws.cell(1, 2 + i, name)
    for r, x in enumerate(xs):
        ws.cell(2 + r, 1, x)
        for i, (_, vals) in enumerate(series):
            ws.cell(2 + r, 2 + i, vals[r])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def cache(vals, numeric):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
    if numeric:
        return f'<c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numCache>'
    return f'<c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache>'


def series_xml(kind, i, name, vals):
    n = len(vals)
    xs = SCATTER_X if kind == "scatter" else CATS
    tx = f'<c:tx><c:strRef><c:f>Sheet1!${col(1 + i)}$1</c:f>{cache([name], False)}</c:strRef></c:tx>'
    xref = f"Sheet1!$A$2:$A${n + 1}"
    vref = f"Sheet1!${col(1 + i)}$2:${col(1 + i)}${n + 1}"
    if kind == "scatter":
        data = (f'<c:xVal><c:numRef><c:f>{xref}</c:f>{cache(xs, True)}</c:numRef></c:xVal>'
                f'<c:yVal><c:numRef><c:f>{vref}</c:f>{cache(vals, True)}</c:numRef></c:yVal>'
                '<c:smooth val="0"/>')
    else:
        data = (f'<c:cat><c:strRef><c:f>{xref}</c:f>{cache(xs, False)}</c:strRef></c:cat>'
                f'<c:val><c:numRef><c:f>{vref}</c:f>{cache(vals, True)}</c:numRef></c:val>')
    extra = '<c:marker><c:symbol val="none"/></c:marker>' if kind == "line" else ""
    if kind == "scatter":
        extra = '<c:marker><c:symbol val="circle"/></c:marker>'
    return f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>{tx}{extra}{data}</c:ser>'


AXES = ('<c:catAx><c:axId val="100"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
        '<c:axPos val="b"/><c:crossAx val="200"/></c:catAx>'
        '<c:valAx><c:axId val="200"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
        '<c:axPos val="l"/><c:majorGridlines/><c:crossAx val="100"/></c:valAx>')
SCATTER_AXES = AXES.replace("<c:catAx>", "<c:valAx>").replace("</c:catAx>", "</c:valAx>", 1)


def chart_xml(kind, series, embedded):
    sers = "".join(series_xml(kind, i, n, v) for i, (n, v) in enumerate(series))
    if kind == "bar":
        plot = (f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>{sers}'
                '<c:gapWidth val="219"/><c:overlap val="-27"/><c:axId val="100"/><c:axId val="200"/></c:barChart>' + AXES)
    elif kind == "line":
        plot = (f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{sers}'
                '<c:marker val="1"/><c:axId val="100"/><c:axId val="200"/></c:lineChart>' + AXES)
    elif kind == "pie":
        plot = f'<c:pieChart><c:varyColors val="1"/>{sers}<c:firstSliceAng val="0"/></c:pieChart>'
    else:
        plot = (f'<c:scatterChart><c:scatterStyle val="lineMarker"/><c:varyColors val="0"/>{sers}'
                '<c:axId val="100"/><c:axId val="200"/></c:scatterChart>' + SCATTER_AXES)
    ext = '<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData>' if embedded else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<c:chartSpace xmlns:c="{C}" xmlns:a="{A}" xmlns:r="{R}"><c:roundedCorners val="0"/>'
            f'<c:chart><c:autoTitleDeleted val="0"/><c:plotArea><c:layout/>{plot}</c:plotArea>'
            f'<c:legend><c:legendPos val="b"/><c:overlay val="0"/></c:legend><c:plotVisOnly val="1"/></c:chart>'
            f'{ext}</c:chartSpace>')


def frame(n):
    return (f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{10 + n}" name="Chart {n}"/><p:cNvGraphicFramePr/>'
            f'<p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="{80 * EMU}" y="{60 * EMU}"/>'
            f'<a:ext cx="{560 * EMU}" cy="{360 * EMU}"/></p:xfrm><a:graphic>'
            f'<a:graphicData uri="{C}"><c:chart xmlns:c="{C}" r:id="rIdC{n}"/></a:graphicData></a:graphic></p:graphicFrame>')


src, dst = sys.argv[1], sys.argv[2]
zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
names = set(zin.namelist())
for item in zin.infolist():
    data = zin.read(item.filename)
    m = re.match(r"ppt/slides/(_rels/)?slide(\d+)\.xml(\.rels)?$", item.filename)
    if m and int(m.group(2)) <= len(CHARTS):
        n = int(m.group(2))
        xml = data.decode()
        if m.group(1):
            xml = xml.replace("</Relationships>",
                              f'<Relationship Id="rIdC{n}" Type="{R}/chart" Target="../charts/chart{n}.xml"/></Relationships>')
        else:
            xml = xml.replace("</p:spTree>", frame(n) + "</p:spTree>", 1)
        data = xml.encode()
    if item.filename == "[Content_Types].xml":
        xml = data.decode()
        over = "".join(f'<Override PartName="/ppt/charts/chart{n}.xml" '
                       'ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
                       for n in range(1, len(CHARTS) + 1))
        if 'Extension="xlsx"' not in xml:
            over += ('<Default Extension="xlsx" '
                     'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>')
        data = xml.replace("</Types>", over + "</Types>").encode()
    zout.writestr(item, data)
for n, (kind, series, embedded) in enumerate(CHARTS, 1):
    zout.writestr(f"ppt/charts/chart{n}.xml", chart_xml(kind, series, embedded))
    if embedded:
        zout.writestr(f"ppt/charts/_rels/chart{n}.xml.rels",
                      '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                      '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                      f'<Relationship Id="rId1" Type="{R}/package" '
                      f'Target="../embeddings/Microsoft_Excel_Worksheet{n}.xlsx"/></Relationships>')
        zout.writestr(f"ppt/embeddings/Microsoft_Excel_Worksheet{n}.xlsx", workbook(kind, series))
zout.close()
